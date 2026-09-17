# -*- coding: utf-8 -*-
"""
AI 智慧校對（Claude API）
------------------------
每篇投稿存檔後，在背景靜默呼叫 Claude 做：
  - 錯別字修正、語意潤飾
  - 橋牌術語規範化（叫品、合約、防禦信號等）
產出「校對後全文」＋「修訂對照表（原句 / 修訂句 / 修正理由）」。

設計原則：
  - 對投稿者完全透明／無感 —— 任何錯誤都只寫進伺服器 log（logging 模組），
    絕不呼叫 st.warning / st.error，不會出現在投稿頁面上。
  - 未設定 st.secrets[anthropic] 時 `is_enabled()` 回傳 False，呼叫端略過。

需要的 st.secrets 結構：
    [anthropic]
    api_key = "sk-ant-...."
    model = "claude-sonnet-5"   # 選填，預設 claude-sonnet-5
"""

from __future__ import annotations

import json
import logging
import re
from datetime import datetime

import streamlit as st

log = logging.getLogger("bridge_portal.proofreader")

_DEFAULT_MODEL = "claude-sonnet-5"
_MAX_TOKENS = 8000

_SYSTEM_PROMPT = """你是《中橋季刊》的專業編輯校對員，精通繁體中文寫作與橋牌（Bridge / 康橋）術語。

任務：校對投稿者提供的文章全文，包括：
1. 錯別字、標點、語法修正
2. 語意潤飾，使行文流暢、合乎雜誌出版水準，但保留原作者的語氣與觀點，不改變原意
3. 橋牌術語規範化：統一叫品（如「1NT」「2♠」）、合約、莊家／防禦方、防禦信號、
   梅花／方塊／紅心／黑桃（♣♦♥♠）等用語的寫法，修正常見誤用
4. 不要增刪與原文無關的內容，不要自行擴寫或大幅改寫段落結構

請只回傳一個 JSON 物件（不要有任何其他文字、不要用 Markdown code fence），格式如下：
{
  "corrected_text": "修訂後的完整文章全文（保留原本的分段換行）",
  "revisions": [
    {"original": "原句", "revised": "修訂句", "reason": "修正理由（簡短一句話）"}
  ]
}

若原文完全沒有需要修改之處，"revisions" 給空陣列 []，但 "corrected_text" 仍要填入原文。"""


def _conf() -> dict | None:
    try:
        if "anthropic" not in st.secrets:
            return None
        c = dict(st.secrets["anthropic"])
    except Exception:  # noqa: BLE001
        return None
    if not c.get("api_key"):
        return None
    c.setdefault("model", _DEFAULT_MODEL)
    return c


def is_enabled() -> bool:
    return _conf() is not None


def status_text() -> str:
    c = _conf()
    if not c:
        return "未設定"
    return f"已設定（{c['model']}）"


def _extract_json(text: str) -> dict | None:
    """盡量從模型回覆中挖出合法 JSON；容許前後有多餘文字或 code fence。"""
    text = text.strip()
    try:
        return json.loads(text)
    except Exception:  # noqa: BLE001
        pass
    fence = re.search(r"```(?:json)?\s*(\{.*\})\s*```", text, re.DOTALL)
    if fence:
        try:
            return json.loads(fence.group(1))
        except Exception:  # noqa: BLE001
            pass
    brace = re.search(r"\{.*\}", text, re.DOTALL)
    if brace:
        try:
            return json.loads(brace.group(0))
        except Exception:  # noqa: BLE001
            pass
    return None


def proofread(content: str, title: str, category: str) -> dict | None:
    """
    呼叫 Claude 校對文章。成功回傳：
        {"corrected_text": str, "revisions": [{"original","revised","reason"}...],
         "model": str, "checked_at": ISO 時間字串}
    任何失敗（未設定金鑰、API 錯誤、JSON 解析失敗…）一律回傳 None，僅寫 log，
    絕不中斷或影響投稿流程。
    """
    conf = _conf()
    if not conf or not (content or "").strip():
        return None

    try:
        import anthropic

        client = anthropic.Anthropic(api_key=conf["api_key"])
        user_prompt = (
            f"文章標題：{title}\n"
            f"建議刊登單元：{category}\n"
            f"---\n"
            f"{content.strip()}"
        )
        resp = client.messages.create(
            model=conf["model"],
            max_tokens=_MAX_TOKENS,
            system=_SYSTEM_PROMPT,
            messages=[{"role": "user", "content": user_prompt}],
        )
        raw = "".join(
            block.text for block in resp.content if getattr(block, "type", "") == "text"
        )
        data = _extract_json(raw)
        if not data or "corrected_text" not in data:
            log.warning("proofreader: 無法解析 Claude 回覆為 JSON，原始回覆前 300 字：%s", raw[:300])
            return None

        revisions = data.get("revisions") or []
        cleaned_revisions = [
            {
                "original": str(r.get("original", "")).strip(),
                "revised": str(r.get("revised", "")).strip(),
                "reason": str(r.get("reason", "")).strip(),
            }
            for r in revisions
            if isinstance(r, dict)
        ]
        return {
            "corrected_text": str(data["corrected_text"]).strip(),
            "revisions": cleaned_revisions,
            "model": conf["model"],
            "checked_at": datetime.now().isoformat(timespec="seconds"),
        }
    except Exception as exc:  # noqa: BLE001
        log.warning("proofreader: 校對失敗（%s）：%s", title, exc)
        return None


def build_revision_docx(meta: dict, proof: dict) -> bytes:
    """產出「修訂對照表」Word 文件（原句／修訂句／理由 表格 ＋ 校對後全文）。"""
    import io

    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()

    title = doc.add_heading(f"《{meta.get('title', '')}》AI 校對對照表", level=1)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0xBA, 0x18, 0x1B)

    info = doc.add_paragraph()
    info.add_run(
        f"投稿人：{meta.get('submitter_name', '')}　"
        f"期別：{meta.get('year', '')} {meta.get('season', '')}　"
        f"建議單元：{meta.get('category', '')}\n"
        f"校對時間：{proof.get('checked_at', '')}　模型：{proof.get('model', '')}"
    ).font.size = Pt(10)

    doc.add_heading("修訂對照表", level=2)
    revisions = proof.get("revisions") or []
    if not revisions:
        doc.add_paragraph("（本篇未發現需要修訂之處）")
    else:
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 2"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "原句", "修訂句", "修正理由"
        for r in revisions:
            row = table.add_row().cells
            row[0].text = r.get("original", "")
            row[1].text = r.get("revised", "")
            row[2].text = r.get("reason", "")

    doc.add_heading("校對後全文", level=2)
    for para in (proof.get("corrected_text") or "").split("\n"):
        doc.add_paragraph(para)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
